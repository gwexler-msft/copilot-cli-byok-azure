from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "presentations" / "gov-byok-briefing.pptx"

NAVY = RGBColor(18, 35, 52)
INK = RGBColor(35, 48, 58)
MUTED = RGBColor(91, 107, 119)
WHITE = RGBColor(255, 255, 255)
PAPER = RGBColor(246, 248, 249)
CYAN = RGBColor(0, 120, 140)
CYAN_LIGHT = RGBColor(218, 241, 244)
GOLD = RGBColor(230, 174, 45)
GOLD_LIGHT = RGBColor(252, 242, 210)
GREEN = RGBColor(42, 132, 102)
GREEN_LIGHT = RGBColor(222, 242, 234)
RED = RGBColor(176, 67, 63)
LINE = RGBColor(198, 207, 213)


def set_background(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, margin=0.04,
             valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def add_title(slide, number, title, kicker=None):
    add_text(slide, f"{number:02d}", 0.58, 0.38, 0.55, 0.34, 13, CYAN, True)
    if kicker:
        add_text(slide, kicker.upper(), 1.12, 0.38, 4.8, 0.28, 10, MUTED, True)
    add_text(slide, title, 0.58, 0.78, 12.0, 0.72, 28, NAVY, True, font="Aptos Display")
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.62), Inches(1.47), Inches(1.0), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = GOLD
    line.line.fill.background()


def add_footer(slide, label="PRIVATE COPILOT | CUSTOMER-CONTROLLED AI"):
    add_text(slide, label, 0.62, 7.18, 6.2, 0.18, 8, MUTED, True)
    add_text(slide, "AZURE GOVERNMENT + COMMERCIAL", 10.08, 7.18, 2.62, 0.18, 8, MUTED, True,
             align=PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, title, body, accent=CYAN, fill=WHITE,
             title_size=16, body_size=12):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = LINE
    shape.line.width = Pt(0.8)
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    add_text(slide, title, x + 0.24, y + 0.20, w - 0.42, 0.36, title_size, NAVY, True)
    add_text(slide, body, x + 0.24, y + 0.69, w - 0.42, h - 0.86, body_size, MUTED)
    return shape


def add_pill(slide, text, x, y, w, color=CYAN, fill=CYAN_LIGHT):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.36)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    add_text(slide, text, x, y + 0.02, w, 0.26, 10, color, True, align=PP_ALIGN.CENTER)


def add_notes(slide, notes):
    frame = slide.notes_slide.notes_text_frame
    frame.text = notes


def add_bullets(slide, items, x, y, w, h, size=16, color=INK, spacing=10):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(spacing)
        paragraph.text = f"•  {item}"
    return box


def add_arrow(slide, x1, y1, x2, y2, color=CYAN, width=2.0):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    connector.line.end_arrowhead = True
    return connector


def build_deck():
    presentation = Presentation()
    presentation.core_properties.author = ""
    presentation.core_properties.last_modified_by = ""
    presentation.core_properties.comments = ""
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    # Slide 1
    slide = presentation.slides.add_slide(blank)
    set_background(slide, NAVY)
    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = GOLD
    band.line.fill.background()
    add_pill(slide, "AZURE GOVERNMENT + COMMERCIAL", 0.82, 0.72, 2.65, GOLD, GOLD_LIGHT)
    add_text(slide, "Private Copilot,\nCustomer-Controlled AI", 0.78, 1.45, 8.1, 1.75,
             38, WHITE, True, font="Aptos Display")
    add_text(slide, "A reusable BYOK gateway that connects Copilot CLI and VS Code to private\nMicrosoft Foundry through an internal Azure API Management AI gateway.",
             0.82, 3.55, 7.7, 1.02, 18, RGBColor(210, 222, 229))
    for index, (label, color) in enumerate([
        ("PRIVATE DATA PLANE", CYAN),
        ("MANAGED IDENTITY", GREEN),
        ("GOVERNED CONSUMPTION", GOLD),
    ]):
        add_pill(slide, label, 0.82 + index * 2.45, 5.13, 2.17, color,
                 RGBColor(31, 58, 72) if color != GOLD else RGBColor(67, 59, 33))
    # Abstract trust-boundary visual
    for index, (label, fill) in enumerate([
        ("COPILOT", RGBColor(34, 69, 83)),
        ("APIM", RGBColor(0, 102, 120)),
        ("FOUNDRY", RGBColor(34, 107, 83)),
    ]):
        x = 9.25 + (index % 2) * 1.75
        y = 1.62 + index * 1.55
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.HEXAGON, Inches(x), Inches(y), Inches(1.65), Inches(1.0)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = RGBColor(99, 130, 144)
        add_text(slide, label, x, y + 0.32, 1.65, 0.25, 12, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(slide, "PRESENTATION BRIEF | SEPTEMBER 2026", 0.82, 6.91, 4.2, 0.2, 9,
             RGBColor(147, 169, 180), True)
    add_notes(slide, "This is not a replacement developer tool. It is a controlled model path behind the tools developers already use. The work packages the architecture as reusable, parameterized infrastructure rather than a one-off deployment.")

    # Slide 2
    slide = presentation.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, 2, "The customer problem we solved", "Why this work matters")
    add_card(slide, 0.62, 1.86, 3.78, 3.55, "Developer experience",
             "Teams want AI-assisted coding in Copilot CLI and VS Code without learning a separate tool or workflow.", CYAN)
    add_card(slide, 4.78, 1.86, 3.78, 3.55, "Regulated trust boundary",
             "Security needs a private model data plane, controlled identities, centralized policy, and a defensible traffic path.", RED)
    add_card(slide, 8.94, 1.86, 3.78, 3.55, "Repeatable delivery",
             "Platform teams need one maintainable implementation across Commercial and Government, not a bespoke rebuild per customer.", GREEN)
    answer = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.72), Inches(5.82), Inches(9.9), Inches(0.82)
    )
    answer.fill.solid()
    answer.fill.fore_color.rgb = NAVY
    answer.line.fill.background()
    add_text(slide, "THE ANSWER", 1.98, 6.07, 1.15, 0.23, 10, GOLD, True)
    add_text(slide, "Internal APIM AI gateway → private Microsoft Foundry", 3.12, 5.98, 7.95, 0.36,
             18, WHITE, True)
    add_footer(slide)
    add_notes(slide, "The important distinction is customer control of the inference path. GitHub SaaS is not the default model path. APIM validates the caller, enforces policy, removes the inbound credential, and authenticates separately to the model backend.")

    # Slide 3
    slide = presentation.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, 3, "How a request flows", "Architecture")
    nodes = [
        (0.58, 2.05, 2.25, "1  DEVELOPER", "Copilot CLI\nor VS Code", CYAN_LIGHT, CYAN),
        (3.35, 2.05, 2.45, "2  PRIVATE ACCESS", "P2S VPN /\nin-VNet client", GOLD_LIGHT, GOLD),
        (6.32, 1.75, 2.72, "3  AI GATEWAY", "Internal APIM\nauth • policy • route", RGBColor(225, 234, 240), NAVY),
        (9.58, 2.05, 2.65, "4  PRIVATE MODEL", "Foundry via PE\nMI authentication", GREEN_LIGHT, GREEN),
    ]
    for x, y, w, heading, body, fill, accent in nodes:
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.55)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = accent
        shape.line.width = Pt(1.2)
        add_text(slide, heading, x + 0.18, y + 0.20, w - 0.36, 0.25, 11, accent, True)
        add_text(slide, body, x + 0.18, y + 0.64, w - 0.36, 0.66, 16, NAVY, True)
    add_arrow(slide, 2.83, 2.82, 3.35, 2.82)
    add_arrow(slide, 5.80, 2.82, 6.32, 2.82)
    add_arrow(slide, 9.04, 2.82, 9.58, 2.82)
    add_card(slide, 4.62, 4.43, 4.1, 1.37, "OBSERVE + ATTRIBUTE",
             "Log Analytics / App Insights\nper developer • per model • throttles", CYAN, WHITE, 13, 11)
    add_arrow(slide, 7.68, 3.30, 7.02, 4.43, CYAN, 1.4)
    add_pill(slide, "CALLER CREDENTIAL STOPS AT APIM", 0.68, 6.22, 3.08, RED, RGBColor(247, 229, 227))
    add_pill(slide, "PUBLIC MODEL ACCESS OFF", 4.06, 6.22, 2.72, GREEN, GREEN_LIGHT)
    add_pill(slide, "BACKEND KEYS OFF", 7.08, 6.22, 2.15, GREEN, GREEN_LIGHT)
    add_pill(slide, "POLICY IN ONE PLACE", 9.54, 6.22, 2.54, CYAN, CYAN_LIGHT)
    add_footer(slide)
    add_notes(slide, "The inbound and backend identities are deliberately separate. A developer's APIM key or JWT never reaches Foundry. The backend sees the APIM managed identity. Chat Completions, Responses, model discovery, streaming, and optional Anthropic-wire clients are handled at the gateway.")

    # Slide 4
    slide = presentation.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, 4, "Government is a profile, not a fork", "Sovereign-cloud design")
    add_card(slide, 0.62, 1.84, 5.82, 4.78, "SAME DESIGN + CODEBASE",
             "• One subscription-scope Bicep/azd implementation\n\n• Cloud profiles select authorities, audiences, DNS, regions, and endpoints\n\n• Same policy, operations, and deployment model\n\n• Default inference path stays private inside the Government tenant",
             CYAN, WHITE, 16, 14)
    add_card(slide, 6.84, 1.84, 5.87, 4.78, "GOV-SPECIFIC ENGINEERING",
             "• Classic APIM Developer for pilot; Premium for production\n\n• Sovereign Entra and Cognitive Services endpoints\n\n• VNet-injected self-hosted runners validate the private path\n\n• Workspace tables and portal/in-VNet telemetry checks\n\n• API-version and Responses routing proven against Gov behavior",
             GOLD, WHITE, 16, 14)
    add_pill(slide, "NO GOVERNMENT-SPECIFIC SOURCE BRANCH", 4.31, 6.73, 4.70, NAVY,
             RGBColor(226, 232, 236))
    add_footer(slide)
    add_notes(slide, "Azure Government is not treated as Commercial with a different URL. The template parameterizes authorities, audiences, DNS, and availability. The default Gov inference path remains inside the Government tenant and private network. A cross-cloud Commercial model backend exists only as an explicit opt-in and changes that boundary, so it must be disclosed and approved separately.")

    # Slide 5
    slide = presentation.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, 5, "What we delivered and proved", "Pilot evidence")
    milestones = [
        ("01", "PRIVATE PLATFORM", "Internal APIM, private endpoints, managed identity, public/local auth disabled"),
        ("02", "CLIENT COVERAGE", "Copilot CLI + VS Code; Chat Completions, Responses, streaming, model discovery"),
        ("03", "GOVERNANCE", "Per-developer tiers, limits, routing, throttles, and attributable telemetry"),
        ("04", "OPERATIONS", "Private onboarding, VNet CI runners, dual-cloud smoke tests, documented runbooks"),
    ]
    for index, (number, heading, body) in enumerate(milestones):
        y = 1.72 + index * 1.23
        add_text(slide, number, 0.73, y + 0.15, 0.62, 0.36, 16, CYAN, True)
        add_text(slide, heading, 1.48, y + 0.06, 2.45, 0.28, 12, NAVY, True)
        add_text(slide, body, 3.67, y, 7.75, 0.65, 14, MUTED)
        if index < 3:
            divider = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1.48), Inches(y + 0.88), Inches(10.55), Inches(0.012)
            )
            divider.fill.solid()
            divider.fill.fore_color.rgb = LINE
            divider.line.fill.background()
    badge = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.HEXAGON, Inches(11.48), Inches(2.35), Inches(1.22), Inches(1.12)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = GREEN
    badge.line.fill.background()
    add_text(slide, "2.2", 11.48, 2.60, 1.22, 0.35, 20, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(slide, "RELEASE", 11.48, 3.03, 1.22, 0.20, 8, WHITE, True, align=PP_ALIGN.CENTER)
    add_pill(slide, "LIVE-VALIDATED ON COMMERCIAL + GOVERNMENT PILOTS", 3.47, 6.67, 6.45,
             GREEN, GREEN_LIGHT)
    add_footer(slide)
    add_notes(slide, "Say pilot validated, not universally production certified. The strongest proof is the automated dual-cloud regression loop: the same source provisions both dev environments, tests from inside each private VNet, and preserves long-lived pilots for demos and canary checks.")

    # Slide 6
    slide = presentation.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, 6, "The difficult parts are now reusable IP", "Lessons de-risked")
    lessons = [
        ("FLEET AUTH", "Long-lived APIM subscription keys avoid hourly client refresh failures; JWT remains an option."),
        ("PROTOCOL", "Policies handle streaming usage, reasoning parameters, Responses paths, and client metadata."),
        ("SAFETY", "A coding-specific content policy reduces false jailbreak blocks without discarding controls."),
        ("GOV TELEMETRY", "Workspace tables and in-VNet assertions account for sovereign query-plane differences."),
        ("PRIVATE CI", "Runner bootstrap, Key Vault rotation, ordering, and smoke validation are automated."),
        ("DAY-2 OPS", "Failure signatures and recovery steps are captured in a symptom-driven runbook."),
    ]
    for index, (heading, body) in enumerate(lessons):
        column = index % 2
        row = index // 2
        x = 0.62 + column * 6.17
        y = 1.75 + row * 1.55
        add_card(slide, x, y, 5.72, 1.24, heading, body,
                 CYAN if column == 0 else GOLD, WHITE, 13, 11)
    callout = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(2.12), Inches(6.39), Inches(9.10), Inches(0.53)
    )
    callout.fill.solid()
    callout.fill.fore_color.rgb = NAVY
    callout.line.fill.background()
    add_text(slide, "Configure and deploy. Do not rediscover the platform constraints per engagement.",
             2.12, 6.53, 9.10, 0.22, 13, WHITE, True, align=PP_ALIGN.CENTER)
    add_footer(slide)
    add_notes(slide, "This slide is the reusable-IP argument. Much of the value is not the resource diagram; it is the set of tested policy fixes and operating procedures that prevent every customer team from rediscovering the same platform constraints.")

    # Slide 7
    slide = presentation.slides.add_slide(blank)
    set_background(slide, NAVY)
    add_text(slide, "07", 0.62, 0.42, 0.5, 0.25, 12, GOLD, True)
    add_text(slide, "From pilot to customer platform", 0.62, 0.87, 7.5, 0.68, 30, WHITE, True,
             font="Aptos Display")
    add_text(slide, "A controlled adoption path with measurable customer outcomes", 0.66, 1.58, 7.2, 0.36,
             15, RGBColor(195, 211, 220))
    outcomes = [
        "Private inference and a customer-controlled trust boundary",
        "Managed identity instead of backend keys on developer machines",
        "Per-developer governance, attribution, and model choice",
        "One operating model across Commercial and Government",
        "Reusable delivery assets that shorten time to pilot",
    ]
    add_bullets(slide, outcomes, 0.65, 2.30, 6.10, 3.45, 15, WHITE, 13)
    steps = [
        ("1", "PROFILE", "Cloud, region, models, network, APIM tier"),
        ("2", "IDENTITY", "Subscription key or JWT by fleet need"),
        ("3", "PROVE", "CI deploy + private smoke + telemetry"),
        ("4", "SCALE", "Cohort rollout, tune, promote tier"),
    ]
    for index, (number, heading, body) in enumerate(steps):
        y = 2.13 + index * 1.08
        circle = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, Inches(7.48), Inches(y), Inches(0.48), Inches(0.48)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = GOLD
        circle.line.fill.background()
        add_text(slide, number, 7.48, y + 0.11, 0.48, 0.20, 11, NAVY, True, align=PP_ALIGN.CENTER)
        add_text(slide, heading, 8.17, y - 0.01, 1.48, 0.24, 11, GOLD, True)
        add_text(slide, body, 9.44, y - 0.01, 3.18, 0.55, 12, WHITE)
    add_text(slide, "NEXT DECISION", 7.48, 6.56, 1.40, 0.20, 9, GOLD, True)
    add_text(slide, "Select a customer pilot cohort and production target tier.", 8.90, 6.48, 3.72, 0.42,
             13, WHITE, True)
    add_notes(slide, "Close on the decision this enables: a customer can begin with a tightly scoped pilot, prove the private trust boundary and developer workflow, then scale using the same artifacts and operating model. Production should use an SLA-backed APIM tier and customer-specific security, capacity, and compliance review.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    return presentation


if __name__ == "__main__":
    deck = build_deck()
    print(f"Created {OUTPUT} with {len(deck.slides)} slides")